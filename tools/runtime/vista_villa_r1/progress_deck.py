"""Build an editable local progress deck from actual native evidence."""
import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches,Pt


def main():
    p=argparse.ArgumentParser();p.add_argument('--native',type=Path,required=True)
    p.add_argument('--character-review',type=Path)
    p.add_argument('--concepts',type=Path,required=True);p.add_argument('--out',type=Path,required=True);a=p.parse_args()
    if a.out.exists():raise RuntimeError('Fresh deck directory required')
    proof=json.loads((a.native/'proof.json').read_text())
    if not proof.get('demo_completed'):raise RuntimeError('Use a completed native demonstration')
    a.out.mkdir(parents=True)
    deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
    ink=(29,47,42);muted=(91,103,91);accent=(111,136,108);paper=(244,241,231)
    def text(s,value,x,y,w,h,size=22,color=ink,bold=False):
        tf=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame;tf.word_wrap=True
        for i,line in enumerate(value.split('\n')):
            q=tf.paragraphs[0] if i==0 else tf.add_paragraph();q.text=line;q.font.name='Microsoft JhengHei'
            q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=RGBColor(*color);q.space_after=Pt(12)
    def slide(title,kicker='VISTA  /  VILLA R1'):
        s=deck.slides.add_slide(deck.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(*paper)
        text(s,kicker,.55,.28,12,.4,12,accent,True);text(s,title,.55,.88,12,1.05,30,ink,True)
        text(s,f'2026-09-10  ·  原生 UE 畫面與可操作進度   /   {len(deck.slides)}',.55,7.05,12,.24,10,muted)
        return s
    def photo(s,path,x,y,w,h):
        pic=s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w))
        ratio=pic.width/pic.height;target=w/h
        if ratio>target:
            pic.width=Inches(w);pic.height=Inches(h);crop=(1-target/ratio)/2;pic.crop_left=pic.crop_right=crop
        else:
            pic.width=Inches(w);pic.height=Inches(h);crop=(1-ratio/target)/2;pic.crop_top=pic.crop_bottom=crop
    s=slide('把設計圖做成可走動、可操作的住宅')
    photo(s,a.native/'09_architecture.png',.55,2,8.1,4.55)
    text(s,'兩層別墅\n第一／第三人稱\n抓取、倒水、放回\n保留原本 Home R5',9.05,2.22,3.75,3.8,24)
    s=slide('用同一組設計方向驗收空間')
    photo(s,a.concepts/'03-kitchen.png',.55,2.03,5.95,3.96);photo(s,a.native/'10_architecture.png',6.82,2.03,5.95,3.96)
    text(s,'目標圖：生成的室內設計',.55,6.15,5.95,.6,18);text(s,'目前成果：Unreal 實機畫面',6.82,6.15,5.95,.6,18)
    s=slide('一份尺寸配置，連接所有房間')
    text(s,'一樓',.65,2.03,5.8,.6,26,accent,True);text(s,'客廳、餐廳、廚房、洗衣房\n6.4 m 挑高客廳\n木作、整片石材、落地玻璃',.65,2.85,5.9,2.6,24)
    text(s,'二樓',7,2.03,5.8,.6,26,accent,True);text(s,'兩間臥室、書房、浴室\n3.2 m 樓層高度\n20 階樓梯與迴廊防護欄',7,2.85,5.6,2.6,24)
    text(s,'建築外框 16 × 12 m；樓梯以角色移動與實際碰撞驗證。',.65,6.02,12,.7,20,muted)
    if a.character_review:
        s=slide('人物與動作：沿用資產，逐項改善')
        photo(s,a.character_review/'face.png',.55,2.0,3.65,3.65)
        photo(s,a.character_review/'body.png',4.45,2.0,3.0,4.25)
        text(s,'保留 53 根骨骼\n同步第一／第三人稱身體\n眼睛、角膜與綁定髮絲\n三段 CMU 步行來源\n203 個姿勢樣本＋IK',7.85,2.05,4.8,3.95,23)
        text(s,'左圖為 Blender 資產檢查；非 Unreal 實機。臉部、服裝及動作庫仍需精修。',.55,6.45,12.2,.45,15,muted)
    s=slide('現在可以示範的連續操作')
    shots=[('01_empty_hands.png','空手看見雙手'),('01b_look_down_body.png','低頭看見身體'),('03_grasped_carafe.png','抓取玻璃壺'),
        ('04_pouring.png','對準杯子倒水'),('05_placed.png','放回支撐面'),('08_stair_landing.png','走上二樓')]
    for i,(file,label) in enumerate(shots):
        x=.55+(i%3)*4.15;y=1.96+(i//3)*2.42
        photo(s,a.native/file,x,y,3.96,1.84);text(s,label,x,y+1.88,3.96,.36,16)
    s=slide('液體：可見水位、可控制的來源')
    photo(s,a.native/'04_pouring.png',.55,2.05,6.65,3.94)
    text(s,f'600 ml 初始壺水\n本次倒入杯中：{proof["receiver_ml"]:.1f} ml\n水量帳本殘差：{proof["mass_residual_ml"]:.2g} ml\n水源關閉後持續模擬',7.58,2.1,5.2,3.4,23)
    text(s,'混合模型：守恆帳本＋容器液面＋彈道細水柱＋粗網格 FLIP；並非已量測的 FLIP 質量。',.55,6.32,12.25,.55,16,muted)
    s=slide('把 VISTA 與 EgoArgus 延伸到可觀察後果')
    text(s,'VISTA\n情境生成與稽核',.65,2.25,3.5,1.4,26,accent,True)
    text(s,'EgoArgus\n畫面／對話證據關係',4.75,2.25,3.9,1.4,26,accent,True)
    text(s,'3D 執行環境\n動作、介入與後果',9.03,2.25,3.8,1.4,26,accent,True)
    text(s,'下一個實驗：固定起始世界，比較沉默、及時提醒與延遲提醒造成的結果。\n研究設計與文獻已記錄；完整助理閉環與世界狀態還原仍需接入。',.65,4.44,12,1.8,23)
    s=slide('驗證結果與仍待改善的部分')
    median=proof.get('frame_ms_median',0);p95=proof.get('frame_ms_p95',0)
    text(s,f'原生操作完成；可上樓\n{proof["motion_library_frames"]} 個動捕姿勢樣本\n核心、啟動與舊版回歸檢查通過\nFrame time 中位數 {median:.1f} ms / p95 {p95:.1f} ms',.65,2.08,6.15,3.72,22)
    text(s,'尚未達到完整 GTA 品質\n需要更多高品質動作與人物細節\n需提升流體精度及即時效能\n完整研究閉環尚未完成',7.23,2.08,5.3,3.72,22,muted)
    text(s,'效能為該次 1920 × 1080 GPU 操作流程，包含自動截圖；不宣稱穩定 30 fps。',.65,6.31,12,.55,16,muted)
    s=slide('Demo 操作')
    text(s,'在 Moonlight / Sunshine 選擇「VISTA Villa R1」',.65,2.05,12,.85,27,accent,True)
    text(s,'WASD ＋滑鼠：走動與觀看\nTab：第一／第三人稱\nE：拿取或放回\nP：拿壺時對杯子倒水',.65,3.25,6.15,2.8,24)
    text(s,'F：操作附近水龍頭\nH：完整示範流程\nR：重設物件與水量\n原本「VISTA Home R5」仍保留',7.3,3.25,5.4,2.8,24)
    for s in deck.slides:s.notes_slide.notes_text_frame.text='Native evidence: '+str(a.native)+'\nConcept source: '+str(a.concepts)+'\nGenerated reference images are not native render evidence.'
    dest=a.out/'VISTA-Villa-progress.zh-TW.pptx';deck.save(dest)
    (a.out/'manifest.json').write_text(json.dumps({'slides':len(deck.slides),'native_proof':str(a.native/'proof.json'),'file':str(dest)},indent=2)+'\n')
    print(dest)


if __name__=='__main__':main()
